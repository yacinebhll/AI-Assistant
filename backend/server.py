from fastapi import FastAPI, APIRouter, HTTPException, Depends, status, UploadFile, File, Form, Query
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict, EmailStr
from typing import List, Optional
import uuid
from datetime import datetime, timezone, timedelta
import bcrypt
import jwt
from emergentintegrations.llm.chat import LlmChat, UserMessage
import aiofiles
import io
import re

# File processing imports
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch

ROOT_DIR = Path(__file__).parent
UPLOAD_DIR = ROOT_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# JWT Config
JWT_SECRET = os.environ.get('JWT_SECRET', 'medical_ai_secret_key')
JWT_ALGORITHM = "HS256"
JWT_EXPIRATION_HOURS = 24

# LLM Key
EMERGENT_LLM_KEY = os.environ.get('EMERGENT_LLM_KEY')

# Create the main app
app = FastAPI(title="Assistant IA Médical API")

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# Security
security = HTTPBearer()

# Logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# Allowed file types
ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.doc', '.xlsx', '.xls', '.png', '.jpg', '.jpeg', '.gif', '.txt'}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB

# ============== MODELS ==============

class UserRole:
    ADMIN = "admin"
    DIRECTION = "direction"
    PERSONNEL_SOIGNANT = "personnel_soignant"

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserResponse(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    email: str
    name: str
    role: str
    created_at: str

class UserRoleUpdate(BaseModel):
    role: str

class DocumentResponse(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    title: str
    content: str
    category_id: str
    category_name: Optional[str] = None
    visibility: List[str]
    created_by: str
    created_by_name: Optional[str] = None
    created_at: str
    updated_at: str
    file_name: Optional[str] = None
    file_type: Optional[str] = None
    file_url: Optional[str] = None
    is_favorite: Optional[bool] = False

class CategoryCreate(BaseModel):
    name: str
    description: Optional[str] = ""

class CategoryResponse(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    name: str
    description: str
    created_at: str

class AIQuestionRequest(BaseModel):
    question: str

class AIResponse(BaseModel):
    response: str
    sources: Optional[List[str]] = []

class DashboardStats(BaseModel):
    total_documents: int
    total_users: int
    total_favorites: int
    unread_notifications: int
    documents_by_category: List[dict]
    users_by_role: List[dict]
    recent_documents: List[dict]

class NotificationResponse(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    user_id: str
    document_id: str
    document_title: str
    message: str
    is_read: bool
    created_at: str

class FavoriteResponse(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    user_id: str
    document_id: str
    created_at: str

class ExportPDFRequest(BaseModel):
    document_ids: List[str]

# ============== AUTH HELPERS ==============

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')

def verify_password(password: str, hashed: str) -> bool:
    return bcrypt.checkpw(password.encode('utf-8'), hashed.encode('utf-8'))

def create_token(user_id: str, email: str, role: str) -> str:
    payload = {
        "user_id": user_id,
        "email": email,
        "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRATION_HOURS)
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

def decode_token(token: str) -> dict:
    try:
        return jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expiré")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token invalide")

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    token_data = decode_token(credentials.credentials)
    user = await db.users.find_one({"id": token_data["user_id"]}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=401, detail="Utilisateur non trouvé")
    return user

async def require_admin(current_user: dict = Depends(get_current_user)):
    if current_user["role"] != UserRole.ADMIN:
        raise HTTPException(status_code=403, detail="Accès réservé aux administrateurs")
    return current_user

async def require_admin_or_direction(current_user: dict = Depends(get_current_user)):
    if current_user["role"] not in [UserRole.ADMIN, UserRole.DIRECTION]:
        raise HTTPException(status_code=403, detail="Accès réservé à l'administration et direction")
    return current_user

# ============== FILE PROCESSING HELPERS ==============

def extract_text_from_pdf(file_path: Path) -> str:
    try:
        reader = PdfReader(str(file_path))
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text.strip()
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return ""

def extract_text_from_docx(file_path: Path) -> str:
    try:
        doc = DocxDocument(str(file_path))
        text = "\n".join([para.text for para in doc.paragraphs])
        return text.strip()
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return ""

def extract_text_from_xlsx(file_path: Path) -> str:
    try:
        wb = load_workbook(str(file_path), data_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) if cell else "" for cell in row])
                if row_text.strip():
                    text.append(row_text)
        return "\n".join(text).strip()
    except Exception as e:
        logger.error(f"XLSX extraction error: {e}")
        return ""

def extract_text_from_file(file_path: Path, file_type: str) -> str:
    if file_type == '.pdf':
        return extract_text_from_pdf(file_path)
    elif file_type in ['.docx', '.doc']:
        return extract_text_from_docx(file_path)
    elif file_type in ['.xlsx', '.xls']:
        return extract_text_from_xlsx(file_path)
    elif file_type == '.txt':
        try:
            return file_path.read_text(encoding='utf-8')
        except:
            return ""
    return ""

async def create_notification_for_users(document_id: str, document_title: str, visibility: List[str], created_by: str):
    """Create notifications for users who can see the document"""
    now = datetime.now(timezone.utc).isoformat()
    
    # Get all users with matching roles except the creator
    users = await db.users.find(
        {"role": {"$in": visibility}, "id": {"$ne": created_by}},
        {"_id": 0, "id": 1}
    ).to_list(1000)
    
    notifications = []
    for user in users:
        notifications.append({
            "id": str(uuid.uuid4()),
            "user_id": user["id"],
            "document_id": document_id,
            "document_title": document_title,
            "message": f"Nouveau document: {document_title}",
            "is_read": False,
            "created_at": now
        })
    
    if notifications:
        await db.notifications.insert_many(notifications)

# ============== AUTH ENDPOINTS ==============

@api_router.post("/auth/register", response_model=dict)
async def register(user_data: UserCreate):
    existing = await db.users.find_one({"email": user_data.email})
    if existing:
        raise HTTPException(status_code=400, detail="Email déjà utilisé")
    
    user_count = await db.users.count_documents({})
    role = UserRole.ADMIN if user_count == 0 else UserRole.PERSONNEL_SOIGNANT
    
    user_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    
    user_doc = {
        "id": user_id,
        "email": user_data.email,
        "name": user_data.name,
        "password": hash_password(user_data.password),
        "role": role,
        "created_at": now
    }
    
    await db.users.insert_one(user_doc)
    token = create_token(user_id, user_data.email, role)
    
    return {
        "token": token,
        "user": {
            "id": user_id,
            "email": user_data.email,
            "name": user_data.name,
            "role": role
        }
    }

@api_router.post("/auth/login", response_model=dict)
async def login(login_data: UserLogin):
    user = await db.users.find_one({"email": login_data.email}, {"_id": 0})
    if not user or not verify_password(login_data.password, user["password"]):
        raise HTTPException(status_code=401, detail="Email ou mot de passe incorrect")
    
    token = create_token(user["id"], user["email"], user["role"])
    
    return {
        "token": token,
        "user": {
            "id": user["id"],
            "email": user["email"],
            "name": user["name"],
            "role": user["role"]
        }
    }

@api_router.get("/auth/me", response_model=UserResponse)
async def get_me(current_user: dict = Depends(get_current_user)):
    return UserResponse(
        id=current_user["id"],
        email=current_user["email"],
        name=current_user["name"],
        role=current_user["role"],
        created_at=current_user["created_at"]
    )

# ============== USERS ENDPOINTS ==============

@api_router.get("/users", response_model=List[UserResponse])
async def get_users(current_user: dict = Depends(require_admin)):
    users = await db.users.find({}, {"_id": 0, "password": 0}).to_list(1000)
    return [UserResponse(**u) for u in users]

@api_router.put("/users/{user_id}/role", response_model=UserResponse)
async def update_user_role(user_id: str, role_data: UserRoleUpdate, current_user: dict = Depends(require_admin)):
    if role_data.role not in [UserRole.ADMIN, UserRole.DIRECTION, UserRole.PERSONNEL_SOIGNANT]:
        raise HTTPException(status_code=400, detail="Rôle invalide")
    
    result = await db.users.update_one({"id": user_id}, {"$set": {"role": role_data.role}})
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Utilisateur non trouvé")
    
    user = await db.users.find_one({"id": user_id}, {"_id": 0, "password": 0})
    return UserResponse(**user)

@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, current_user: dict = Depends(require_admin)):
    if user_id == current_user["id"]:
        raise HTTPException(status_code=400, detail="Vous ne pouvez pas supprimer votre propre compte")
    result = await db.users.delete_one({"id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Utilisateur non trouvé")
    return {"message": "Utilisateur supprimé"}

# ============== CATEGORIES ENDPOINTS ==============

@api_router.post("/document-categories", response_model=CategoryResponse)
async def create_category(category: CategoryCreate, current_user: dict = Depends(require_admin)):
    category_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    
    doc = {
        "id": category_id,
        "name": category.name,
        "description": category.description or "",
        "created_at": now
    }
    
    await db.categories.insert_one(doc)
    return CategoryResponse(**doc)

@api_router.get("/document-categories", response_model=List[CategoryResponse])
async def get_categories(current_user: dict = Depends(get_current_user)):
    categories = await db.categories.find({}, {"_id": 0}).to_list(1000)
    return [CategoryResponse(**c) for c in categories]

@api_router.delete("/document-categories/{category_id}")
async def delete_category(category_id: str, current_user: dict = Depends(require_admin)):
    doc_count = await db.documents.count_documents({"category_id": category_id})
    if doc_count > 0:
        raise HTTPException(status_code=400, detail=f"Cette catégorie contient {doc_count} document(s). Supprimez-les d'abord.")
    
    result = await db.categories.delete_one({"id": category_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Catégorie non trouvée")
    return {"message": "Catégorie supprimée"}

# ============== DOCUMENTS ENDPOINTS ==============

@api_router.post("/documents", response_model=DocumentResponse)
async def create_document(
    title: str = Form(...),
    content: str = Form(""),
    category_id: str = Form(...),
    visibility: str = Form("admin,direction,personnel_soignant"),
    file: Optional[UploadFile] = File(None),
    current_user: dict = Depends(get_current_user)
):
    # Verify category exists
    category = await db.categories.find_one({"id": category_id}, {"_id": 0})
    if not category:
        raise HTTPException(status_code=400, detail="Catégorie non trouvée")
    
    visibility_list = [v.strip() for v in visibility.split(",")]
    
    doc_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    file_name = None
    file_type = None
    file_url = None
    extracted_content = content
    
    # Handle file upload
    if file and file.filename:
        file_ext = Path(file.filename).suffix.lower()
        if file_ext not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=400, detail=f"Type de fichier non autorisé. Types acceptés: {', '.join(ALLOWED_EXTENSIONS)}")
        
        # Read and check file size
        file_content = await file.read()
        if len(file_content) > MAX_FILE_SIZE:
            raise HTTPException(status_code=400, detail="Le fichier dépasse la taille maximale de 10 MB")
        
        # Save file
        file_name = f"{doc_id}{file_ext}"
        file_path = UPLOAD_DIR / file_name
        async with aiofiles.open(file_path, 'wb') as f:
            await f.write(file_content)
        
        file_type = file_ext
        file_url = f"/api/documents/{doc_id}/file"
        
        # Extract text for full-text search
        if file_ext in ['.pdf', '.docx', '.doc', '.xlsx', '.xls', '.txt']:
            extracted_text = extract_text_from_file(file_path, file_ext)
            if extracted_text and not content:
                extracted_content = extracted_text
            elif extracted_text:
                extracted_content = f"{content}\n\n--- Contenu extrait du fichier ---\n{extracted_text}"
    
    doc = {
        "id": doc_id,
        "title": title,
        "content": extracted_content,
        "category_id": category_id,
        "visibility": visibility_list,
        "created_by": current_user["id"],
        "created_at": now,
        "updated_at": now,
        "file_name": file_name,
        "file_type": file_type,
        "file_url": file_url,
        "search_text": f"{title} {extracted_content}".lower()
    }
    
    await db.documents.insert_one(doc)
    
    # Create notifications for other users
    await create_notification_for_users(doc_id, title, visibility_list, current_user["id"])
    
    return DocumentResponse(
        **{k: v for k, v in doc.items() if k != "search_text"},
        category_name=category["name"],
        created_by_name=current_user["name"],
        is_favorite=False
    )

@api_router.get("/documents", response_model=List[DocumentResponse])
async def get_documents(
    category_id: Optional[str] = None,
    search: Optional[str] = None,
    favorites_only: bool = False,
    current_user: dict = Depends(get_current_user)
):
    query = {"visibility": current_user["role"]}
    
    if category_id:
        query["category_id"] = category_id
    
    if search:
        search_lower = search.lower()
        query["$or"] = [
            {"title": {"$regex": search_lower, "$options": "i"}},
            {"content": {"$regex": search_lower, "$options": "i"}},
            {"search_text": {"$regex": search_lower, "$options": "i"}}
        ]
    
    documents = await db.documents.find(query, {"_id": 0, "search_text": 0}).sort("created_at", -1).to_list(1000)
    
    # Get category names and creator names
    category_ids = list(set(d["category_id"] for d in documents))
    creator_ids = list(set(d["created_by"] for d in documents))
    
    categories = await db.categories.find({"id": {"$in": category_ids}}, {"_id": 0}).to_list(100)
    cat_map = {c["id"]: c["name"] for c in categories}
    
    creators = await db.users.find({"id": {"$in": creator_ids}}, {"_id": 0, "id": 1, "name": 1}).to_list(100)
    creator_map = {c["id"]: c["name"] for c in creators}
    
    # Get user's favorites
    favorites = await db.favorites.find({"user_id": current_user["id"]}, {"_id": 0, "document_id": 1}).to_list(1000)
    favorite_ids = set(f["document_id"] for f in favorites)
    
    # Filter favorites if requested
    if favorites_only:
        documents = [d for d in documents if d["id"] in favorite_ids]
    
    return [
        DocumentResponse(
            **d,
            category_name=cat_map.get(d["category_id"], ""),
            created_by_name=creator_map.get(d["created_by"], ""),
            is_favorite=d["id"] in favorite_ids
        ) for d in documents
    ]

@api_router.get("/documents/{doc_id}", response_model=DocumentResponse)
async def get_document(doc_id: str, current_user: dict = Depends(get_current_user)):
    doc = await db.documents.find_one({"id": doc_id, "visibility": current_user["role"]}, {"_id": 0, "search_text": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Document non trouvé")
    
    category = await db.categories.find_one({"id": doc["category_id"]}, {"_id": 0})
    creator = await db.users.find_one({"id": doc["created_by"]}, {"_id": 0, "name": 1})
    
    favorite = await db.favorites.find_one({"user_id": current_user["id"], "document_id": doc_id})
    
    return DocumentResponse(
        **doc,
        category_name=category["name"] if category else "",
        created_by_name=creator["name"] if creator else "",
        is_favorite=favorite is not None
    )

@api_router.get("/documents/{doc_id}/file")
async def get_document_file(doc_id: str, current_user: dict = Depends(get_current_user)):
    doc = await db.documents.find_one({"id": doc_id, "visibility": current_user["role"]}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Document non trouvé")
    
    if not doc.get("file_name"):
        raise HTTPException(status_code=404, detail="Pas de fichier associé")
    
    file_path = UPLOAD_DIR / doc["file_name"]
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Fichier non trouvé")
    
    # Determine content type
    ext = Path(doc["file_name"]).suffix.lower()
    content_types = {
        '.pdf': 'application/pdf',
        '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        '.doc': 'application/msword',
        '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        '.xls': 'application/vnd.ms-excel',
        '.png': 'image/png',
        '.jpg': 'image/jpeg',
        '.jpeg': 'image/jpeg',
        '.gif': 'image/gif',
        '.txt': 'text/plain'
    }
    
    async def file_iterator():
        async with aiofiles.open(file_path, 'rb') as f:
            while chunk := await f.read(8192):
                yield chunk
    
    return StreamingResponse(
        file_iterator(),
        media_type=content_types.get(ext, 'application/octet-stream'),
        headers={"Content-Disposition": f"inline; filename={doc['file_name']}"}
    )

@api_router.put("/documents/{doc_id}", response_model=DocumentResponse)
async def update_document(
    doc_id: str,
    title: Optional[str] = Form(None),
    content: Optional[str] = Form(None),
    category_id: Optional[str] = Form(None),
    visibility: Optional[str] = Form(None),
    file: Optional[UploadFile] = File(None),
    current_user: dict = Depends(get_current_user)
):
    doc = await db.documents.find_one({"id": doc_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Document non trouvé")
    
    # Check permissions: Admin/Direction can edit all, others only their own
    can_edit = (
        current_user["role"] in [UserRole.ADMIN, UserRole.DIRECTION] or
        doc["created_by"] == current_user["id"]
    )
    if not can_edit:
        raise HTTPException(status_code=403, detail="Non autorisé")
    
    update_data = {"updated_at": datetime.now(timezone.utc).isoformat()}
    
    if title is not None:
        update_data["title"] = title
    if content is not None:
        update_data["content"] = content
    if category_id is not None:
        update_data["category_id"] = category_id
    if visibility is not None:
        update_data["visibility"] = [v.strip() for v in visibility.split(",")]
    
    # Handle new file upload
    if file and file.filename:
        file_ext = Path(file.filename).suffix.lower()
        if file_ext not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=400, detail="Type de fichier non autorisé")
        
        file_content = await file.read()
        if len(file_content) > MAX_FILE_SIZE:
            raise HTTPException(status_code=400, detail="Fichier trop volumineux")
        
        # Delete old file
        if doc.get("file_name"):
            old_file = UPLOAD_DIR / doc["file_name"]
            if old_file.exists():
                old_file.unlink()
        
        # Save new file
        file_name = f"{doc_id}{file_ext}"
        file_path = UPLOAD_DIR / file_name
        async with aiofiles.open(file_path, 'wb') as f:
            await f.write(file_content)
        
        update_data["file_name"] = file_name
        update_data["file_type"] = file_ext
        update_data["file_url"] = f"/api/documents/{doc_id}/file"
        
        # Extract text
        if file_ext in ['.pdf', '.docx', '.doc', '.xlsx', '.xls', '.txt']:
            extracted_text = extract_text_from_file(file_path, file_ext)
            if extracted_text:
                current_content = update_data.get("content", doc.get("content", ""))
                update_data["content"] = f"{current_content}\n\n--- Contenu extrait ---\n{extracted_text}"
    
    # Update search text
    final_title = update_data.get("title", doc["title"])
    final_content = update_data.get("content", doc["content"])
    update_data["search_text"] = f"{final_title} {final_content}".lower()
    
    await db.documents.update_one({"id": doc_id}, {"$set": update_data})
    
    updated_doc = await db.documents.find_one({"id": doc_id}, {"_id": 0, "search_text": 0})
    category = await db.categories.find_one({"id": updated_doc["category_id"]}, {"_id": 0})
    creator = await db.users.find_one({"id": updated_doc["created_by"]}, {"_id": 0, "name": 1})
    favorite = await db.favorites.find_one({"user_id": current_user["id"], "document_id": doc_id})
    
    return DocumentResponse(
        **updated_doc,
        category_name=category["name"] if category else "",
        created_by_name=creator["name"] if creator else "",
        is_favorite=favorite is not None
    )

@api_router.delete("/documents/{doc_id}")
async def delete_document(doc_id: str, current_user: dict = Depends(get_current_user)):
    doc = await db.documents.find_one({"id": doc_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Document non trouvé")
    
    # Check permissions
    can_delete = (
        current_user["role"] in [UserRole.ADMIN, UserRole.DIRECTION] or
        doc["created_by"] == current_user["id"]
    )
    if not can_delete:
        raise HTTPException(status_code=403, detail="Non autorisé")
    
    # Delete file
    if doc.get("file_name"):
        file_path = UPLOAD_DIR / doc["file_name"]
        if file_path.exists():
            file_path.unlink()
    
    # Delete related data
    await db.documents.delete_one({"id": doc_id})
    await db.favorites.delete_many({"document_id": doc_id})
    await db.notifications.delete_many({"document_id": doc_id})
    
    return {"message": "Document supprimé"}

# ============== EXPORT PDF ENDPOINT ==============

@api_router.post("/documents/export-pdf")
async def export_documents_pdf(request: ExportPDFRequest, current_user: dict = Depends(get_current_user)):
    if not request.document_ids:
        raise HTTPException(status_code=400, detail="Aucun document sélectionné")
    
    # Get documents
    documents = await db.documents.find(
        {"id": {"$in": request.document_ids}, "visibility": current_user["role"]},
        {"_id": 0, "search_text": 0}
    ).to_list(100)
    
    if not documents:
        raise HTTPException(status_code=404, detail="Aucun document trouvé")
    
    # Get categories
    cat_ids = list(set(d["category_id"] for d in documents))
    categories = await db.categories.find({"id": {"$in": cat_ids}}, {"_id": 0}).to_list(100)
    cat_map = {c["id"]: c["name"] for c in categories}
    
    # Create PDF
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=72)
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18, spaceAfter=12)
    subtitle_style = ParagraphStyle('Subtitle', parent=styles['Normal'], fontSize=10, textColor='gray', spaceAfter=20)
    heading_style = ParagraphStyle('Heading', parent=styles['Heading2'], fontSize=14, spaceBefore=20, spaceAfter=10)
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=11, spaceAfter=12, leading=14)
    
    story = []
    
    # Header
    story.append(Paragraph("Assistant IA Médical - Export de documents", title_style))
    story.append(Paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", subtitle_style))
    story.append(Spacer(1, 0.3*inch))
    
    for doc_data in documents:
        story.append(Paragraph(doc_data["title"], heading_style))
        story.append(Paragraph(f"Catégorie: {cat_map.get(doc_data['category_id'], 'N/A')} | Date: {doc_data['created_at'][:10]}", subtitle_style))
        
        # Clean and format content
        content = doc_data["content"].replace("\n", "<br/>")
        story.append(Paragraph(content, body_style))
        story.append(Spacer(1, 0.3*inch))
    
    doc.build(story)
    buffer.seek(0)
    
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=documents_export.pdf"}
    )

# ============== FAVORITES ENDPOINTS ==============

@api_router.post("/favorites/{document_id}")
async def add_favorite(document_id: str, current_user: dict = Depends(get_current_user)):
    # Check document exists and user can access it
    doc = await db.documents.find_one({"id": document_id, "visibility": current_user["role"]})
    if not doc:
        raise HTTPException(status_code=404, detail="Document non trouvé")
    
    # Check if already favorited
    existing = await db.favorites.find_one({"user_id": current_user["id"], "document_id": document_id})
    if existing:
        return {"message": "Déjà en favoris"}
    
    favorite = {
        "id": str(uuid.uuid4()),
        "user_id": current_user["id"],
        "document_id": document_id,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    await db.favorites.insert_one(favorite)
    return {"message": "Ajouté aux favoris", "id": favorite["id"]}

@api_router.delete("/favorites/{document_id}")
async def remove_favorite(document_id: str, current_user: dict = Depends(get_current_user)):
    result = await db.favorites.delete_one({"user_id": current_user["id"], "document_id": document_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Favori non trouvé")
    return {"message": "Retiré des favoris"}

@api_router.get("/favorites", response_model=List[DocumentResponse])
async def get_favorites(current_user: dict = Depends(get_current_user)):
    # Get user's favorites
    favorites = await db.favorites.find({"user_id": current_user["id"]}, {"_id": 0}).to_list(1000)
    doc_ids = [f["document_id"] for f in favorites]
    
    if not doc_ids:
        return []
    
    # Get documents
    documents = await db.documents.find(
        {"id": {"$in": doc_ids}, "visibility": current_user["role"]},
        {"_id": 0, "search_text": 0}
    ).to_list(1000)
    
    # Get category and creator names
    cat_ids = list(set(d["category_id"] for d in documents))
    creator_ids = list(set(d["created_by"] for d in documents))
    
    categories = await db.categories.find({"id": {"$in": cat_ids}}, {"_id": 0}).to_list(100)
    cat_map = {c["id"]: c["name"] for c in categories}
    
    creators = await db.users.find({"id": {"$in": creator_ids}}, {"_id": 0, "id": 1, "name": 1}).to_list(100)
    creator_map = {c["id"]: c["name"] for c in creators}
    
    return [
        DocumentResponse(
            **d,
            category_name=cat_map.get(d["category_id"], ""),
            created_by_name=creator_map.get(d["created_by"], ""),
            is_favorite=True
        ) for d in documents
    ]

# ============== NOTIFICATIONS ENDPOINTS ==============

@api_router.get("/notifications", response_model=List[NotificationResponse])
async def get_notifications(current_user: dict = Depends(get_current_user)):
    notifications = await db.notifications.find(
        {"user_id": current_user["id"]},
        {"_id": 0}
    ).sort("created_at", -1).limit(50).to_list(50)
    
    return [NotificationResponse(**n) for n in notifications]

@api_router.get("/notifications/unread-count")
async def get_unread_count(current_user: dict = Depends(get_current_user)):
    count = await db.notifications.count_documents({"user_id": current_user["id"], "is_read": False})
    return {"count": count}

@api_router.put("/notifications/{notification_id}/read")
async def mark_notification_read(notification_id: str, current_user: dict = Depends(get_current_user)):
    result = await db.notifications.update_one(
        {"id": notification_id, "user_id": current_user["id"]},
        {"$set": {"is_read": True}}
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Notification non trouvée")
    return {"message": "Marquée comme lue"}

@api_router.put("/notifications/read-all")
async def mark_all_notifications_read(current_user: dict = Depends(get_current_user)):
    await db.notifications.update_many(
        {"user_id": current_user["id"], "is_read": False},
        {"$set": {"is_read": True}}
    )
    return {"message": "Toutes les notifications marquées comme lues"}

# ============== AI ASSISTANT ENDPOINTS ==============

async def get_documents_for_ai(user_role: str, limit: int = 10) -> List[dict]:
    docs = await db.documents.find(
        {"visibility": user_role},
        {"_id": 0, "title": 1, "content": 1, "category_id": 1}
    ).sort("created_at", -1).limit(limit).to_list(limit)
    return docs

@api_router.get("/ai-assistant/suggestions-context", response_model=AIResponse)
async def get_context_suggestions(current_user: dict = Depends(get_current_user)):
    docs = await get_documents_for_ai(current_user["role"], 5)
    
    if not docs:
        return AIResponse(response="Aucun document disponible pour générer des suggestions.", sources=[])
    
    context = "\n\n".join([f"Document: {d['title']}\nContenu: {d['content'][:500]}" for d in docs])
    
    try:
        chat = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"context-{current_user['id']}-{uuid.uuid4()}",
            system_message="Tu es un assistant médical IA. Tu analyses les documents et fournis des suggestions contextuelles utiles pour le personnel médical. Réponds toujours en français."
        ).with_model("openai", "gpt-5.2")
        
        message = UserMessage(text=f"Analyse ces documents récents et fournis 3-5 suggestions pratiques basées sur leur contenu:\n\n{context}")
        response = await chat.send_message(message)
        
        return AIResponse(response=response, sources=[d["title"] for d in docs])
    except Exception as e:
        logger.error(f"AI Error: {e}")
        return AIResponse(response="Erreur lors de la génération des suggestions. Veuillez réessayer.", sources=[])

@api_router.get("/ai-assistant/suggestions-enhanced", response_model=AIResponse)
async def get_enhanced_suggestions(current_user: dict = Depends(get_current_user)):
    docs = await get_documents_for_ai(current_user["role"], 10)
    
    if not docs:
        return AIResponse(response="Aucun document disponible pour générer des suggestions avancées.", sources=[])
    
    context = "\n\n".join([f"### {d['title']}\n{d['content']}" for d in docs])
    
    try:
        chat = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"enhanced-{current_user['id']}-{uuid.uuid4()}",
            system_message="Tu es un assistant médical IA expert. Tu fournis des analyses approfondies avec points clés, résumés et recommandations. Réponds toujours en français de manière structurée."
        ).with_model("openai", "gpt-5.2")
        
        message = UserMessage(text=f"""Analyse ces documents et fournis:
1. Un résumé exécutif (3-4 lignes)
2. Les points clés importants (liste)
3. Des recommandations pratiques
4. Des alertes ou points d'attention si nécessaire

Documents:
{context}""")
        response = await chat.send_message(message)
        
        return AIResponse(response=response, sources=[d["title"] for d in docs])
    except Exception as e:
        logger.error(f"AI Error: {e}")
        return AIResponse(response="Erreur lors de la génération des suggestions avancées. Veuillez réessayer.", sources=[])

@api_router.post("/ai-assistant/qa", response_model=AIResponse)
async def ask_question(request: AIQuestionRequest, current_user: dict = Depends(get_current_user)):
    docs = await get_documents_for_ai(current_user["role"], 15)
    
    if not docs:
        return AIResponse(response="Aucun document disponible pour répondre à votre question.", sources=[])
    
    context = "\n\n".join([f"### {d['title']}\n{d['content']}" for d in docs])
    
    try:
        chat = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=f"qa-{current_user['id']}-{uuid.uuid4()}",
            system_message="""Tu es un assistant médical IA spécialisé dans l'analyse de documents administratifs et médicaux.
Tu réponds aux questions en te basant UNIQUEMENT sur les documents fournis.
Si la réponse n'est pas dans les documents, dis-le clairement.
Réponds toujours en français de manière professionnelle et précise."""
        ).with_model("openai", "gpt-5.2")
        
        message = UserMessage(text=f"""Documents de référence:
{context}

Question de l'utilisateur: {request.question}

Réponds à cette question en te basant sur les documents ci-dessus.""")
        response = await chat.send_message(message)
        
        await db.qa_history.insert_one({
            "id": str(uuid.uuid4()),
            "user_id": current_user["id"],
            "question": request.question,
            "response": response,
            "created_at": datetime.now(timezone.utc).isoformat()
        })
        
        return AIResponse(response=response, sources=[d["title"] for d in docs])
    except Exception as e:
        logger.error(f"AI Error: {e}")
        return AIResponse(response="Erreur lors du traitement de votre question. Veuillez réessayer.", sources=[])

@api_router.get("/ai-assistant/qa-history", response_model=List[dict])
async def get_qa_history(current_user: dict = Depends(get_current_user)):
    history = await db.qa_history.find(
        {"user_id": current_user["id"]},
        {"_id": 0}
    ).sort("created_at", -1).limit(20).to_list(20)
    return history

# ============== DASHBOARD ENDPOINTS ==============

@api_router.get("/dashboard/stats", response_model=DashboardStats)
async def get_dashboard_stats(current_user: dict = Depends(get_current_user)):
    total_docs = await db.documents.count_documents({"visibility": current_user["role"]})
    
    if current_user["role"] == UserRole.ADMIN:
        total_users = await db.users.count_documents({})
    else:
        total_users = 0
    
    # Favorites count
    total_favorites = await db.favorites.count_documents({"user_id": current_user["id"]})
    
    # Unread notifications
    unread_notifications = await db.notifications.count_documents({"user_id": current_user["id"], "is_read": False})
    
    # Documents by category
    pipeline = [
        {"$match": {"visibility": current_user["role"]}},
        {"$group": {"_id": "$category_id", "count": {"$sum": 1}}}
    ]
    docs_by_cat = await db.documents.aggregate(pipeline).to_list(100)
    
    cat_ids = [d["_id"] for d in docs_by_cat]
    categories = await db.categories.find({"id": {"$in": cat_ids}}, {"_id": 0}).to_list(100)
    cat_map = {c["id"]: c["name"] for c in categories}
    
    docs_by_category = [{"name": cat_map.get(d["_id"], "Sans catégorie"), "count": d["count"]} for d in docs_by_cat]
    
    # Users by role (admin only)
    users_by_role = []
    if current_user["role"] == UserRole.ADMIN:
        role_pipeline = [{"$group": {"_id": "$role", "count": {"$sum": 1}}}]
        users_agg = await db.users.aggregate(role_pipeline).to_list(10)
        role_names = {"admin": "Admin", "direction": "Direction", "personnel_soignant": "Personnel Soignant"}
        users_by_role = [{"name": role_names.get(u["_id"], u["_id"]), "count": u["count"]} for u in users_agg]
    
    # Recent documents
    recent_docs = await db.documents.find(
        {"visibility": current_user["role"]},
        {"_id": 0, "id": 1, "title": 1, "created_at": 1, "category_id": 1}
    ).sort("created_at", -1).limit(5).to_list(5)
    
    recent_documents = [{"id": d["id"], "title": d["title"], "date": d["created_at"], "category": cat_map.get(d.get("category_id"), "")} for d in recent_docs]
    
    return DashboardStats(
        total_documents=total_docs,
        total_users=total_users,
        total_favorites=total_favorites,
        unread_notifications=unread_notifications,
        documents_by_category=docs_by_category,
        users_by_role=users_by_role,
        recent_documents=recent_documents
    )

# ============== SEED DATA ==============

@api_router.post("/seed-data")
async def seed_demo_data(current_user: dict = Depends(require_admin)):
    now = datetime.now(timezone.utc).isoformat()
    
    cat_count = await db.categories.count_documents({})
    if cat_count > 0:
        return {"message": "Données déjà initialisées"}
    
    categories = [
        {"id": str(uuid.uuid4()), "name": "Protocoles Médicaux", "description": "Protocoles et procédures médicales", "created_at": now},
        {"id": str(uuid.uuid4()), "name": "Ressources Humaines", "description": "Documents RH et administratifs", "created_at": now},
        {"id": str(uuid.uuid4()), "name": "Formations", "description": "Supports et programmes de formation", "created_at": now},
        {"id": str(uuid.uuid4()), "name": "Réglementations", "description": "Textes réglementaires et conformité", "created_at": now},
    ]
    
    await db.categories.insert_many(categories)
    
    documents = [
        {
            "id": str(uuid.uuid4()),
            "title": "Protocole d'hygiène des mains",
            "content": """Protocole standardisé pour l'hygiène des mains dans l'établissement.

1. Se mouiller les mains
2. Appliquer le savon
3. Frotter pendant 30 secondes minimum
4. Rincer abondamment
5. Sécher avec une serviette à usage unique

Ce protocole doit être appliqué avant et après chaque contact patient.""",
            "category_id": categories[0]["id"],
            "visibility": [UserRole.ADMIN, UserRole.DIRECTION, UserRole.PERSONNEL_SOIGNANT],
            "created_by": current_user["id"],
            "created_at": now,
            "updated_at": now,
            "search_text": "protocole hygiène mains savon frotter rincer sécher patient"
        },
        {
            "id": str(uuid.uuid4()),
            "title": "Guide d'accueil des nouveaux employés",
            "content": """Guide complet pour l'intégration des nouveaux membres de l'équipe.

• Jour 1: Présentation des locaux et de l'équipe
• Semaine 1: Formation aux outils et procédures
• Mois 1: Évaluation et ajustements

Contact RH: rh@hopital.fr""",
            "category_id": categories[1]["id"],
            "visibility": [UserRole.ADMIN, UserRole.DIRECTION],
            "created_by": current_user["id"],
            "created_at": now,
            "updated_at": now,
            "search_text": "guide accueil nouveaux employés intégration formation rh"
        },
        {
            "id": str(uuid.uuid4()),
            "title": "Formation gestes d'urgence",
            "content": """Programme de formation continue aux gestes d'urgence.

Sessions prévues:
- Réanimation cardio-pulmonaire: Tous les trimestres
- Utilisation du défibrillateur: Mensuel
- Gestion des voies aériennes: Bimensuel

Inscription obligatoire via le portail RH.""",
            "category_id": categories[2]["id"],
            "visibility": [UserRole.ADMIN, UserRole.DIRECTION, UserRole.PERSONNEL_SOIGNANT],
            "created_by": current_user["id"],
            "created_at": now,
            "updated_at": now,
            "search_text": "formation gestes urgence réanimation défibrillateur voies aériennes"
        },
        {
            "id": str(uuid.uuid4()),
            "title": "Conformité RGPD - Données patients",
            "content": """Réglementation sur la protection des données patients.

Points clés:
• Consentement explicite requis
• Droit d'accès et de rectification
• Durée de conservation limitée
• Notification en cas de violation

Référent DPO: dpo@hopital.fr""",
            "category_id": categories[3]["id"],
            "visibility": [UserRole.ADMIN, UserRole.DIRECTION],
            "created_by": current_user["id"],
            "created_at": now,
            "updated_at": now,
            "search_text": "conformité rgpd données patients protection consentement dpo"
        },
    ]
    
    await db.documents.insert_many(documents)
    
    return {"message": "Données de démonstration créées avec succès", "categories": len(categories), "documents": len(documents)}

# ============== ROOT ==============

@api_router.get("/")
async def root():
    return {"message": "API Assistant IA Médical", "version": "2.0.0"}

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
